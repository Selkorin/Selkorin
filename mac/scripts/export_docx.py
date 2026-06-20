#!/usr/bin/env python3
"""
export_docx.py — Generate DOCX from a base64-encoded JSON ReportDocument.

Usage:
    python3 export_docx.py --output /path/to/file.docx --data <base64_json>

Requires: pip3 install python-docx
Exits 0 on success, 1 on failure (caller falls back to RTF).
"""

import sys
import json
import base64
import argparse
from pathlib import Path

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, help="Destination .docx path")
    parser.add_argument("--data",   required=True, help="Base64-encoded JSON ReportDocument")
    args = parser.parse_args()

    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("python-docx not installed. Run: pip3 install python-docx", file=sys.stderr)
        sys.exit(1)

    try:
        json_bytes = base64.b64decode(args.data)
        doc_data   = json.loads(json_bytes)
    except Exception as e:
        print(f"Failed to decode input: {e}", file=sys.stderr)
        sys.exit(1)

    doc = Document()

    # ── Styles ────────────────────────────────────────────────
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    # ── Title ─────────────────────────────────────────────────
    title = doc_data.get("title", "Документ")
    p = doc.add_heading(title, level=0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    subtitle = doc_data.get("subtitle", "")
    if subtitle:
        p2 = doc.add_paragraph(subtitle)
        p2.runs[0].bold = True

    created = doc_data.get("createdAt", "")
    if created:
        doc.add_paragraph(f"Создан: {created[:10]}")

    doc.add_paragraph("")

    # ── Sections ──────────────────────────────────────────────
    for section in doc_data.get("sections", []):
        heading = section.get("heading", "")
        if heading:
            doc.add_heading(heading, level=2)

        body = section.get("body", "")
        if body:
            doc.add_paragraph(body)

        for bullet in section.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")

        table_data = section.get("table")
        if table_data:
            headers = table_data.get("headers", [])
            rows    = table_data.get("rows", [])
            if headers:
                t = doc.add_table(rows=1 + len(rows), cols=len(headers))
                t.style = "Table Grid"
                hdr_row = t.rows[0]
                for i, h in enumerate(headers):
                    cell = hdr_row.cells[i]
                    cell.text = h
                    cell.paragraphs[0].runs[0].bold = True
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row):
                        t.rows[ri + 1].cells[ci].text = str(val)
                doc.add_paragraph("")

    # ── Sources ───────────────────────────────────────────────
    sources = doc_data.get("sources", [])
    if sources:
        doc.add_heading("Источники", level=2)
        for i, src in enumerate(sources):
            t  = src.get("title", "")
            url = src.get("url", "")
            doc.add_paragraph(f"[{i+1}] {t} — {url}")

    # ── Save ──────────────────────────────────────────────────
    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Saved: {out_path}")
    sys.exit(0)


if __name__ == "__main__":
    main()
